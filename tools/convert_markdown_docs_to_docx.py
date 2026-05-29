from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOC_DIRS = [
    ROOT / "case_1_chatbot_prompt_injection" / "docs",
    ROOT / "case_2_tax_audit_risk_scoring" / "docs",
]

WIDE_DOCS = {"expected_findings.md", "expected_findings_check_guide.md"}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def apply_table_geometry(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)


def clean_inline(text: str) -> str:
    text = text.replace("`", "")
    text = text.replace("**", "")
    return text


def split_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [clean_inline(part.strip()) for part in stripped.split("|")]


def is_table_separator(line: str) -> bool:
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def table_widths(col_count: int, wide: bool) -> list[float]:
    total = 9.0 if wide else 6.5
    presets = {
        2: [1.8, total - 1.8],
        3: [1.2, 2.4, total - 3.6],
        4: [0.7, 2.0, 3.1, total - 5.8],
        5: [0.45, 1.45, 2.15, 2.25, total - 6.3],
        6: [0.4, 1.2, 1.75, 1.75, 1.75, total - 6.85],
    }
    if col_count in presets:
        return presets[col_count]
    return [total / col_count] * col_count


def setup_document(doc: Document, title: str, wide: bool) -> None:
    section = doc.sections[0]
    if wide:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
    else:
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color in [
        ("Heading 1", 16, RGBColor(0x2E, 0x74, 0xB5)),
        ("Heading 2", 13, RGBColor(0x2E, 0x74, 0xB5)),
        ("Heading 3", 12, RGBColor(0x1F, 0x4D, 0x78)),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = title
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_title(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(10)
    run = para.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)


def add_code_block(doc: Document, lines: list[str]) -> None:
    for line in lines:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(0.2)
        para.paragraph_format.space_after = Pt(1)
        run = para.add_run(line)
        run.font.name = "Courier New"
        run.font.size = Pt(9)


def add_markdown_table(doc: Document, rows: list[list[str]], wide: bool) -> None:
    if not rows:
        return
    col_count = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    widths = table_widths(col_count, wide)
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            if r_idx == 0:
                set_cell_shading(cell, "F2F4F7")
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.size = Pt(8.5)
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(table, widths)
    doc.add_paragraph()


def convert_markdown(md_path: Path) -> Path:
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    wide = md_path.name in WIDE_DOCS
    output_path = md_path.with_suffix(".docx")

    first_heading = next((line.lstrip("# ").strip() for line in lines if line.startswith("# ")), md_path.stem)
    doc = Document()
    setup_document(doc, first_heading, wide)

    idx = 0
    in_code = False
    code_lines: list[str] = []
    title_written = False
    while idx < len(lines):
        line = lines[idx].rstrip()

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue

        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        if line.strip().startswith("|") and idx + 1 < len(lines) and is_table_separator(lines[idx + 1]):
            table_rows = [split_table_row(line)]
            idx += 2
            while idx < len(lines) and lines[idx].strip().startswith("|"):
                table_rows.append(split_table_row(lines[idx]))
                idx += 1
            add_markdown_table(doc, table_rows, wide)
            continue

        if not line.strip():
            idx += 1
            continue

        if line.startswith("# "):
            heading = clean_inline(line[2:].strip())
            if not title_written:
                add_title(doc, heading)
                title_written = True
            else:
                doc.add_heading(heading, level=1)
        elif line.startswith("## "):
            doc.add_heading(clean_inline(line[3:].strip()), level=1)
        elif line.startswith("### "):
            doc.add_heading(clean_inline(line[4:].strip()), level=2)
        elif line.startswith("#### "):
            doc.add_heading(clean_inline(line[5:].strip()), level=3)
        elif re.match(r"^\s*-\s+", line):
            para = doc.add_paragraph(style="List Bullet")
            para.add_run(clean_inline(re.sub(r"^\s*-\s+", "", line)))
        elif re.match(r"^\s*\d+\.\s+", line):
            para = doc.add_paragraph(style="List Number")
            para.add_run(clean_inline(re.sub(r"^\s*\d+\.\s+", "", line)))
        else:
            doc.add_paragraph(clean_inline(line))
        idx += 1

    doc.core_properties.author = "Codex"
    doc.core_properties.title = first_heading
    doc.save(output_path)
    return output_path


def main() -> None:
    outputs = []
    for doc_dir in DOC_DIRS:
        for md_path in sorted(doc_dir.glob("*.md")):
            outputs.append(convert_markdown(md_path))
    print("\n".join(str(path) for path in outputs))


if __name__ == "__main__":
    main()
